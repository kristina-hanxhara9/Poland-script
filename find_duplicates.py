#!/usr/bin/env python3
"""
Find duplicate NIP codes in a results Excel file.

Checks column A (NIP) for duplicates and creates a report
with duplicate counts and highlighted rows.

Usage:
    python find_duplicates.py resultsmobile.xlsx --output duplicates.xlsx
"""

import argparse
import logging
import sys

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger(__name__)

# Highlight colors for duplicates
YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
RED_FILL = PatternFill(start_color="FF6666", end_color="FF6666", fill_type="solid")


def main():
    parser = argparse.ArgumentParser(description="Find duplicate NIP codes in results file.")
    parser.add_argument("input_file", help="Excel results file")
    parser.add_argument("--output", "-o", default="duplicates.xlsx", help="Output file")
    parser.add_argument("--col", default="A", help="Column letter to check for duplicates (default: A)")
    args = parser.parse_args()

    logger.info(f"Reading {args.input_file}")
    try:
        df = pd.read_excel(args.input_file, header=0)
    except FileNotFoundError:
        logger.error(f"File not found: {args.input_file}")
        sys.exit(1)

    # Resolve column letter to index
    col_letter = args.col.upper().strip()
    col_idx = 0
    for ch in col_letter:
        col_idx = col_idx * 26 + (ord(ch) - ord('A') + 1)
    col_idx -= 1

    check_col = df.columns[col_idx]
    logger.info(f"Checking column {col_letter}: '{check_col}'")
    logger.info(f"Total rows: {len(df)}")

    # Clean values
    values = df[check_col].astype(str).str.strip().str.replace("-", "", regex=False).str.replace(" ", "", regex=False)
    values = values.replace(["nan", "None", ""], pd.NA)

    # Find duplicates
    dup_mask = values.duplicated(keep=False)  # marks ALL occurrences
    dup_df = df[dup_mask].copy()
    dup_df["_dup_value"] = values[dup_mask]

    # Count occurrences of each duplicate value
    value_counts = values[dup_mask].value_counts()

    # ========================================
    # PRINT SUMMARY
    # ========================================
    total = len(df)
    empty = values.isna().sum()
    non_empty = total - empty
    unique = values.dropna().nunique()
    dup_values = len(value_counts)
    dup_rows = dup_mask.sum()

    print("\n" + "=" * 55)
    print("DUPLICATE ANALYSIS")
    print("=" * 55)
    print(f"\nColumn checked: {col_letter} ('{check_col}')")
    print(f"Total rows:     {total}")
    print(f"Empty values:   {empty}")
    print(f"Non-empty:      {non_empty}")
    print(f"Unique values:  {unique}")
    print(f"Duplicate values: {dup_values} (appearing {dup_rows} times total)")

    if dup_values > 0:
        print(f"\n--- DUPLICATES (value: count) ---")
        for val, count in value_counts.items():
            # Find the names for this duplicate
            names = df.loc[values == val].iloc[:, 1].tolist() if len(df.columns) > 1 else []
            names_str = ", ".join(str(n) for n in names[:3])
            if len(names) > 3:
                names_str += f" (+{len(names)-3} more)"
            print(f"  {val}: {count}x  -> {names_str}")
    else:
        print("\nNo duplicates found!")

    print("=" * 55)

    # ========================================
    # SAVE OUTPUT
    # ========================================
    logger.info(f"Saving to {args.output}")

    with pd.ExcelWriter(args.output, engine="openpyxl") as writer:
        # Sheet 1: Duplicate summary
        summary_rows = []
        for val, count in value_counts.items():
            names = df.loc[values == val].iloc[:, 1].tolist() if len(df.columns) > 1 else []
            summary_rows.append({
                "duplicate_value": val,
                "count": count,
                "names": "; ".join(str(n) for n in names),
            })
        df_summary = pd.DataFrame(summary_rows)
        if not df_summary.empty:
            df_summary.to_excel(writer, sheet_name="Duplicate Summary", index=False)
        else:
            pd.DataFrame({"result": ["No duplicates found"]}).to_excel(
                writer, sheet_name="Duplicate Summary", index=False
            )

        # Sheet 2: All duplicate rows (full data)
        if not dup_df.empty:
            dup_df.drop(columns=["_dup_value"], errors="ignore").to_excel(
                writer, sheet_name="Duplicate Rows", index=False
            )

        # Sheet 3: Full data with duplicates highlighted
        df.to_excel(writer, sheet_name="All Data (highlighted)", index=False)

    # Highlight duplicates in the "All Data" sheet
    wb = load_workbook(args.output)
    ws = wb["All Data (highlighted)"]

    # Build set of duplicate values for fast lookup
    dup_set = set(value_counts.index)

    for row_idx in range(2, ws.max_row + 1):  # skip header
        cell = ws.cell(row=row_idx, column=col_idx + 1)
        cell_val = str(cell.value or "").strip().replace("-", "").replace(" ", "")
        if cell_val in dup_set:
            count = value_counts.get(cell_val, 0)
            if count >= 4:
                fill = RED_FILL
            elif count >= 3:
                fill = ORANGE_FILL
            else:
                fill = YELLOW_FILL
            # Highlight entire row
            for col in range(1, ws.max_column + 1):
                ws.cell(row=row_idx, column=col).fill = fill

    # Add a legend
    legend_row = ws.max_row + 3
    ws.cell(row=legend_row, column=1, value="LEGEND:")
    ws.cell(row=legend_row + 1, column=1, value="Yellow = 2 duplicates")
    ws.cell(row=legend_row + 1, column=2).fill = YELLOW_FILL
    ws.cell(row=legend_row + 2, column=1, value="Orange = 3 duplicates")
    ws.cell(row=legend_row + 2, column=2).fill = ORANGE_FILL
    ws.cell(row=legend_row + 3, column=1, value="Red = 4+ duplicates")
    ws.cell(row=legend_row + 3, column=2).fill = RED_FILL

    wb.save(args.output)

    logger.info(f"Done! Saved to {args.output}")
    logger.info(f"  Sheet 'Duplicate Summary': {dup_values} duplicate values with counts")
    logger.info(f"  Sheet 'Duplicate Rows': {dup_rows} rows that are duplicates")
    logger.info(f"  Sheet 'All Data (highlighted)': full data with duplicates color-coded")


if __name__ == "__main__":
    main()
